from pptx import Presentation
from ocr import extract_text_from_image

def extract_text_from_ppt(ppt_path):

    presentation = Presentation(
        ppt_path
    )

    all_text = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):

                text = shape.text.strip()

                if text:

                    all_text.append(
                        text
                    )

            if shape.shape_type == 13:

                try:

                    image = shape.image

                    image_bytes = (
                        image.blob
                    )

                    ocr_text = (
                        extract_text_from_image(
                            image_bytes
                        )
                    )

                    if ocr_text.strip():

                        all_text.append(
                            ocr_text
                        )

                except Exception:

                    pass

    return "\n".join(
        all_text
    )